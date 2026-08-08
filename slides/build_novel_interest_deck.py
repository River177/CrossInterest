from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "slides" / "output" / "generative_novel_interest_prediction.pptx"
PRIOR_WORK_IMAGE = ROOT / "slides" / "assets" / "novel_interest_prompt.png"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

INK = "111111"
MUTED = "666666"
PAPER = "FFFFFF"
WHITE = "FFFFFF"
MINT = "111111"
MINT_DARK = "111111"
INDIGO = "111111"
INDIGO_PALE = "F0F0F0"
CORAL = "111111"
CORAL_PALE = "F0F0F0"
GOLD = "111111"
GOLD_PALE = "F0F0F0"
LINE = "CFCFCF"
SOFT = "F7F7F7"
FONT = "Arial"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_background(slide, color: str = PAPER) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    bold: bool = False,
    font: str = FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    line_spacing: float = 1.0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    runs: list[tuple[str, dict]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for text, style in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.name = style.get("font", FONT)
        run.font.size = Pt(style.get("size", size))
        run.font.bold = style.get("bold", False)
        run.font.italic = style.get("italic", False)
        run.font.color.rgb = rgb(style.get("color", color))
    return box


def add_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str = WHITE,
    line: str | None = None,
    radius: bool = True,
    line_width: float = 1,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(
        shape_type,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = LINE,
    width: float = 1.5,
):
    line = slide.shapes.add_connector(
        1,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    return line


def add_circle(
    slide,
    text: str,
    x: float,
    y: float,
    d: float,
    *,
    fill: str,
    color: str = INK,
    size: float = 18,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(d),
        Inches(d),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = rgb(color)
    return shape


def add_pill(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    *,
    fill: str,
    color: str = INK,
    size: float = 11,
):
    shape = add_rect(slide, x, y, w, 0.34, fill=fill, radius=True)
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    p = frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = rgb(color)
    return shape


def add_slide_header(slide, section: str, title: str, number: int) -> None:
    add_text(
        slide,
        section.upper(),
        0.72,
        0.42,
        3.2,
        0.3,
        size=10,
        color=INDIGO,
        bold=True,
    )
    add_text(slide, title, 0.72, 0.78, 11.9, 0.7, size=29, bold=True)
    add_text(
        slide,
        f"{number:02d}",
        12.12,
        0.47,
        0.5,
        0.3,
        size=10,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )


def add_footer(slide, text: str = "") -> None:
    add_line(slide, 0.72, 7.12, 12.61, 7.12, color=LINE, width=0.8)
    add_text(slide, text, 0.72, 7.18, 11.9, 0.2, size=7.5, color=MUTED)


def add_bullet_list(
    slide,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = INK,
    bullet_color: str = INDIGO,
    gap: float = 0.58,
) -> None:
    for index, item in enumerate(items):
        yy = y + index * gap
        add_circle(
            slide,
            "",
            x,
            yy + 0.11,
            0.12,
            fill=bullet_color,
            color=bullet_color,
            size=1,
        )
        add_text(slide, item, x + 0.28, yy, w - 0.28, gap, size=size, color=color)


def add_stat(
    slide,
    value: str,
    label: str,
    x: float,
    y: float,
    w: float,
    *,
    fill: str,
):
    add_rect(slide, x, y, w, 1.05, fill=fill)
    add_text(slide, value, x + 0.18, y + 0.14, w - 0.36, 0.4, size=23, bold=True)
    add_text(
        slide,
        label,
        x + 0.18,
        y + 0.61,
        w - 0.36,
        0.25,
        size=10,
        color=MUTED,
        bold=True,
    )


def make_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 — Title
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_pill(slide, "RESEARCH FRAMING", 0.8, 0.62, 1.72, fill=SOFT, color=MINT_DARK)
    add_rich_text(
        slide,
        [
            ("Generative\n", {"color": INK, "bold": True, "size": 43}),
            ("Novel Interest", {"color": INK, "bold": True, "size": 43}),
            ("\nPrediction", {"color": INK, "bold": True, "size": 43}),
        ],
        0.8,
        1.45,
        6.5,
        2.5,
    )
    add_text(
        slide,
        "Predicting what users may explore beyond their previously observed interests",
        0.83,
        4.38,
        5.7,
        0.9,
        size=20,
        color=MUTED,
    )

    # Title-slide trajectory.
    add_line(slide, 7.52, 2.12, 11.85, 2.12, color=LINE, width=2.5)
    nodes = [
        ("KNOWN", "Music", 7.38, CORAL_PALE, CORAL),
        ("ADJACENT", "Film", 9.05, GOLD_PALE, INK),
        ("NOVEL", "Travel", 10.74, SOFT, MINT_DARK),
    ]
    for eyebrow, value, x, fill, color in nodes:
        add_circle(slide, "", x, 1.76, 0.72, fill=fill, color=color)
        add_text(
            slide,
            eyebrow,
            x - 0.08,
            2.65,
            0.88,
            0.22,
            size=8,
            color=MUTED,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            value,
            x - 0.2,
            2.92,
            1.12,
            0.35,
            size=15,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        slide,
        "User history",
        7.42,
        4.15,
        2.1,
        0.3,
        size=11,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "→",
        9.55,
        4.0,
        0.5,
        0.5,
        size=27,
        color=MINT,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "New interest",
        10.35,
        4.15,
        2.1,
        0.3,
        size=11,
        color=MINT,
        bold=True,
    )
    add_text(
        slide,
        "Research Group Meeting · August 2026",
        0.83,
        6.72,
        5.2,
        0.28,
        size=10,
        color=MUTED,
        bold=True,
    )

    # 2 — Motivation
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Motivation", "Users do not only consume more of the same.", 2)
    add_text(
        slide,
        "A recommender that only exploits observed preferences can be accurate—and still feel repetitive.",
        0.72,
        1.58,
        8.8,
        0.55,
        size=18,
        color=MUTED,
    )
    add_line(slide, 1.15, 3.07, 11.62, 3.07, color=LINE, width=2.5)
    stages = [
        ("1", "Observed", "Known preferences", 1.15, CORAL_PALE, CORAL),
        ("2", "Adjacent", "Plausible expansion", 5.05, GOLD_PALE, INK),
        ("3", "Novel", "Unseen in history", 9.0, SOFT, MINT_DARK),
    ]
    for number, title, detail, x, fill, color in stages:
        add_circle(slide, number, x, 2.67, 0.8, fill=fill, color=color, size=20)
        add_text(
            slide,
            title,
            x - 0.35,
            3.72,
            1.5,
            0.35,
            size=17,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
        add_text(
            slide,
            detail,
            x - 0.67,
            4.12,
            2.15,
            0.4,
            size=12,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
    reasons = [
        ("Interests evolve", "Preferences shift with context and time."),
        ("Discovery matters", "Useful recommendations can open new content regions."),
        ("Exploration is explicit", "Predict the direction—not merely a different item."),
    ]
    for index, (title, detail) in enumerate(reasons):
        x = 0.72 + index * 4.08
        add_rect(slide, x, 5.04, 3.75, 1.45, fill=WHITE, line=LINE)
        add_text(slide, title, x + 0.22, 5.28, 3.2, 0.3, size=15, bold=True)
        add_text(
            slide,
            detail,
            x + 0.22,
            5.75,
            3.2,
            0.48,
            size=11.5,
            color=MUTED,
        )
    add_footer(slide)

    # 3 — Problem shift
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Problem Shift", "From ranking candidates to generating preference expansion.", 3)
    add_rect(slide, 0.72, 1.64, 5.25, 4.92, fill=WHITE, line=LINE)
    add_pill(slide, "CONVENTIONAL", 1.02, 1.95, 1.38, fill=CORAL_PALE, color=CORAL)
    add_text(slide, "Which item ranks highest?", 1.02, 2.52, 4.3, 0.48, size=23, bold=True)
    candidates = [
        ("A", "Sports article", 3),
        ("B", "Sports article", 1),
        ("C", "Finance article", 2),
    ]
    for index, (letter, label, rank) in enumerate(candidates):
        y = 3.27 + index * 0.82
        add_rect(slide, 1.02, y, 4.6, 0.58, fill=SOFT)
        add_circle(slide, letter, 1.16, y + 0.09, 0.4, fill=WHITE, size=11)
        add_text(slide, label, 1.75, y + 0.12, 2.55, 0.25, size=13)
        add_text(
            slide,
            f"#{rank}",
            4.7,
            y + 0.12,
            0.55,
            0.25,
            size=13,
            color=MUTED,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
    add_text(
        slide,
        "Closed candidate space",
        1.02,
        5.92,
        4.2,
        0.28,
        size=11,
        color=MUTED,
        bold=True,
    )

    add_text(
        slide,
        "→",
        6.18,
        3.46,
        0.95,
        0.8,
        size=39,
        color=INDIGO,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 7.36, 1.64, 5.25, 4.92, fill=WHITE, line=LINE)
    add_pill(slide, "PROPOSED", 7.68, 1.95, 1.18, fill=SOFT, color=MINT_DARK)
    add_text(
        slide,
        "What new interest comes next?",
        7.68,
        2.52,
        4.3,
        0.75,
        size=23,
        color=INK,
        bold=True,
    )
    add_rect(slide, 7.68, 3.58, 4.58, 1.18, fill=SOFT, line=LINE)
    add_text(
        slide,
        "New topic:",
        7.96,
        3.87,
        1.55,
        0.3,
        size=14,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "Courts & Legal Cases",
        9.2,
        3.83,
        2.68,
        0.42,
        size=16,
        color=MINT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_text(
        slide,
        "Open-ended semantic prediction",
        7.68,
        5.92,
        4.2,
        0.28,
        size=11,
        color=MUTED,
        bold=True,
    )
    add_footer(slide)

    # 4 — Task definition
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Task Definition", "A clicked interest is novel when it is absent from clicked history.", 4)
    add_rect(slide, 0.72, 1.65, 7.15, 3.52, fill=WHITE, line=LINE)
    add_text(slide, "Observed history", 1.05, 1.98, 2.4, 0.3, size=11, color=MUTED, bold=True)
    history = [
        "Celebrity & Creator Culture",
        "Music",
        "Nutrition & Healthy Eating",
    ]
    for index, interest in enumerate(history):
        y = 2.48 + index * 0.64
        add_circle(slide, str(index + 1), 1.05, y, 0.38, fill=SOFT, color=INK, size=10)
        add_text(slide, interest, 1.58, y + 0.04, 4.0, 0.25, size=14, color=INK)
    add_line(slide, 5.68, 2.1, 5.68, 4.72, color=LINE, width=1.2)
    add_text(slide, "Target", 6.0, 1.98, 1.1, 0.3, size=11, color=MUTED, bold=True)
    add_rect(slide, 5.99, 2.52, 1.52, 1.28, fill=SOFT, line=LINE)
    add_text(
        slide,
        "Courts &\nLegal Cases",
        6.12,
        2.81,
        1.25,
        0.62,
        size=14,
        color=MINT_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "not in history",
        6.02,
        4.15,
        1.46,
        0.25,
        size=10,
        color=MINT,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 8.2, 1.65, 4.42, 3.52, fill=WHITE, line=LINE)
    add_text(slide, "Formal rule", 8.52, 1.98, 2.2, 0.3, size=11, color=INDIGO, bold=True)
    add_rich_text(
        slide,
        [
            ("I", {"italic": True, "bold": True, "size": 30}),
            ("t", {"italic": True, "bold": True, "size": 18}),
            ("+  ∉  I(H", {"bold": True, "size": 30}),
            ("u", {"italic": True, "bold": True, "size": 18}),
            (")", {"bold": True, "size": 30}),
        ],
        8.54,
        2.54,
        3.7,
        0.72,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        "Current clicked interest\nis absent from historical clicked interests",
        8.77,
        3.52,
        3.25,
        0.8,
        size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_pill(slide, "NO CANDIDATES", 8.54, 4.55, 1.45, fill=CORAL_PALE, color=CORAL)
    add_pill(slide, "SECOND-LEVEL", 10.13, 4.55, 1.52, fill=INDIGO_PALE, color=INDIGO)

    add_stat(slide, "1.49M", "TRAIN SAMPLES", 0.72, 5.55, 2.55, fill=INDIGO_PALE)
    add_stat(slide, "243K", "DEV SAMPLES", 3.5, 5.55, 2.55, fill=CORAL_PALE)
    add_stat(slide, "125", "FINE-GRAINED INTERESTS", 6.28, 5.55, 2.55, fill=SOFT)
    add_stat(slide, "55%", "BEHAVIORS WITH ≥1 NOVEL INTEREST", 9.06, 5.55, 3.56, fill=GOLD_PALE)
    add_footer(slide, "Source: reclassified MINDlarge; novelty is defined only against observed clicked history.")

    # 5 — Two input views
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Input Representations", "Two views of the same user history.", 5)
    add_rect(slide, 0.72, 1.63, 5.72, 4.95, fill=WHITE, line=LINE)
    add_pill(slide, "BEHAVIOR-ONLY", 1.02, 1.95, 1.48, fill=CORAL_PALE, color=CORAL)
    add_text(slide, "Read article semantics", 1.02, 2.5, 4.6, 0.4, size=23, bold=True)
    titles = [
        "Jamie Foxx Sparks Romance Rumors…",
        "A race against time to slow Sonoma fire…",
        "What Is Mirin and Why Should You Use It?",
    ]
    for index, title in enumerate(titles):
        y = 3.17 + index * 0.69
        add_circle(slide, str(index + 1), 1.03, y, 0.35, fill=CORAL_PALE, color=CORAL, size=9)
        add_text(slide, title, 1.53, y + 0.02, 4.42, 0.35, size=12.5)
    add_line(slide, 1.03, 5.46, 5.98, 5.46, color=LINE, width=1)
    add_text(
        slide,
        "Can the model infer latent interests from titles?",
        1.03,
        5.75,
        4.8,
        0.45,
        size=13,
        color=MUTED,
        bold=True,
    )

    add_rect(slide, 6.89, 1.63, 5.72, 4.95, fill=WHITE, line=LINE)
    add_pill(slide, "INTEREST-ONLY", 7.19, 1.95, 1.42, fill=SOFT, color=MINT_DARK)
    add_text(slide, "Read the interest trajectory", 7.19, 2.5, 4.7, 0.4, size=23, color=INK, bold=True)
    interests = [
        "Celebrity & Creator Culture",
        "Floods, Droughts & Wildfires",
        "Cooking & Recipes",
    ]
    for index, interest in enumerate(interests):
        y = 3.17 + index * 0.69
        add_circle(slide, str(index + 1), 7.2, y, 0.35, fill=SOFT, color=INK, size=9)
        add_text(slide, interest, 7.7, y + 0.02, 4.35, 0.35, size=12.5, color=INK)
    add_line(slide, 7.2, 5.46, 12.15, 5.46, color=LINE, width=1)
    add_text(
        slide,
        "Is an explicit trajectory sufficient?",
        7.2,
        5.75,
        4.8,
        0.45,
        size=13,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "Same target · full order retained · repeated interests preserved",
        3.2,
        6.72,
        6.95,
        0.26,
        size=10.5,
        color=INDIGO,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    # 6 — Concrete example
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Concrete Example", "Open-ended generation from one observed history.", 6)
    add_rect(slide, 0.72, 1.62, 12.0, 4.98, fill=WHITE, line=LINE)
    add_pill(slide, "BEHAVIOR-ONLY", 1.02, 1.94, 1.48, fill=CORAL_PALE, color=CORAL)
    add_text(
        slide,
        "The headlines of the news items I recently viewed are listed below:",
        1.03,
        2.5,
        5.2,
        0.35,
        size=13,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "1. Jamie Foxx Sparks Romance Rumors With Dana Caprio After Katie Split",
        1.03,
        3.07,
        5.1,
        1.1,
        size=18,
        color=INK,
        bold=True,
    )
    add_text(
        slide,
        "Predict a new topic category that I am likely to be interested in next.",
        1.03,
        4.55,
        5.2,
        0.7,
        size=13,
        color=MUTED,
    )
    add_line(slide, 6.47, 2.04, 6.47, 6.16, color=LINE, width=1.2)
    add_text(slide, "MODEL OUTPUT", 6.88, 2.51, 2.0, 0.3, size=10, color=MINT, bold=True)
    add_rect(slide, 6.88, 3.03, 4.94, 1.28, fill=SOFT, line=LINE)
    add_text(slide, "New topic:", 7.22, 3.48, 1.5, 0.32, size=15, color=MUTED, bold=True)
    add_text(
        slide,
        "Courts & Legal Cases",
        8.62,
        3.42,
        2.82,
        0.43,
        size=18,
        color=MINT,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_pill(slide, "ABSENT FROM HISTORY", 6.88, 4.77, 1.78, fill=SOFT, color=MINT_DARK)
    add_pill(slide, "NO OPTIONS", 8.86, 4.77, 1.17, fill=CORAL_PALE, color=CORAL)
    add_pill(slide, "SEMANTIC LABEL", 10.23, 4.77, 1.48, fill=INDIGO_PALE, color=INDIGO)
    add_text(
        slide,
        "The prediction target is a category name—not a news ID, rank, or explanation.",
        6.88,
        5.48,
        4.93,
        0.48,
        size=13,
        color=MUTED,
    )
    add_footer(slide, "Real sample: train:40:20:behavior_only")

    # 7 — Our method overview
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Our Method", "End-to-end pipeline: data, training, generation, and evaluation.", 7)
    pipeline = [
        ("1", "MIND data", "News metadata\n+ behavior logs"),
        ("2", "Data processing", "Novel-interest\nsample construction"),
        ("3", "Model training", "SFT + GRPO\nGRPO with CoT"),
        ("4", "Generation", "Ranked novel-\ninterest predictions"),
        ("5", "Evaluation", "Judge + Recall\n+ NDCG"),
    ]
    for index, (number, title, detail) in enumerate(pipeline):
        x = 0.58 + index * 2.57
        add_rect(slide, x, 2.15, 2.05, 2.42, fill=WHITE, line=LINE)
        add_circle(slide, number, x + 0.18, 2.37, 0.44, fill=SOFT, color=INK, size=10)
        add_text(slide, title, x + 0.18, 3.02, 1.68, 0.35, size=16, bold=True)
        add_text(
            slide,
            detail,
            x + 0.18,
            3.56,
            1.68,
            0.62,
            size=12,
            color=MUTED,
            align=PP_ALIGN.CENTER,
        )
        if index < len(pipeline) - 1:
            add_text(
                slide,
                "→",
                x + 2.08,
                3.02,
                0.46,
                0.45,
                size=24,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    add_rect(slide, 1.28, 5.16, 10.77, 0.84, fill=SOFT, line=LINE)
    add_text(
        slide,
        "Two parallel input views",
        1.55,
        5.39,
        2.18,
        0.28,
        size=12,
        color=MUTED,
        bold=True,
    )
    add_text(
        slide,
        "Behavior-only: news titles",
        4.05,
        5.36,
        2.48,
        0.32,
        size=14,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Interest-only: second-level interests",
        7.15,
        5.36,
        3.25,
        0.32,
        size=14,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "The same source sample and target are used across both views.",
        2.65,
        6.38,
        8.0,
        0.35,
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    # 8 — Dataset construction
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Data Construction", "From MIND behavior logs to generative training samples.", 8)
    data_steps = [
        (
            "01",
            "Reclassified MIND",
            "Map every news item to a first-level and a fine-grained second-level interest.",
        ),
        (
            "02",
            "Novel target extraction",
            "Group current clicked news by interest and retain targets absent from clicked history.",
        ),
        (
            "03",
            "Prompt construction",
            "Convert each sample into Behavior-only and Interest-only conversational records.",
        ),
    ]
    for index, (number, title, detail) in enumerate(data_steps):
        x = 0.72 + index * 4.08
        add_rect(slide, x, 1.72, 3.76, 2.45, fill=WHITE, line=LINE)
        add_circle(slide, number, x + 0.22, 1.98, 0.5, fill=SOFT, color=INK, size=10)
        add_text(slide, title, x + 0.22, 2.71, 3.25, 0.38, size=18, bold=True)
        add_text(slide, detail, x + 0.22, 3.27, 3.25, 0.68, size=12.5, color=MUTED)
        if index < 2:
            add_text(
                slide,
                "→",
                x + 3.72,
                2.72,
                0.37,
                0.45,
                size=22,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
    add_rect(slide, 0.72, 4.58, 12.0, 1.35, fill=SOFT, line=LINE)
    funnel = [
        ("2.19M", "TRAIN BEHAVIORS"),
        ("1.22M", "WITH ≥1 NOVEL INTEREST"),
        ("1.49M", "TRAIN SAMPLES"),
        ("243K", "DEV SAMPLES"),
        ("125", "SECOND-LEVEL LABELS"),
    ]
    for index, (value, label) in enumerate(funnel):
        x = 0.93 + index * 2.37
        add_text(slide, value, x, 4.85, 1.9, 0.38, size=22, bold=True, align=PP_ALIGN.CENTER)
        add_text(
            slide,
            label,
            x,
            5.35,
            1.9,
            0.25,
            size=8.5,
            color=MUTED,
            bold=True,
            align=PP_ALIGN.CENTER,
        )
    add_text(
        slide,
        "Novelty rule: target interest ∉ historical clicked-interest set",
        2.3,
        6.38,
        8.7,
        0.34,
        size=13,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Negative impressions and candidate lists are not used.")

    # 9 — Training data examples
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Training Data Examples", "Two input views, one shared novel-interest target.", 9)
    examples = [
        (
            "BEHAVIOR-ONLY",
            "Input: news titles",
            "The headlines of the news items I recently viewed are listed below:\n\n"
            "1. Jamie Foxx Sparks Romance Rumors With Dana Caprio After Katie Split\n\n"
            "Predict a new topic category that I am likely to be interested in next.",
        ),
        (
            "INTEREST-ONLY",
            "Input: second-level interests",
            "The topic categories of the news items I recently viewed are listed below:\n\n"
            "1. Celebrity & Creator Culture\n\n"
            "Predict a new topic category that I am likely to be interested in next.",
        ),
    ]
    for index, (task_type, subtitle, prompt) in enumerate(examples):
        x = 0.72 + index * 6.17
        add_rect(slide, x, 1.62, 5.72, 4.75, fill=WHITE, line=LINE)
        add_pill(slide, task_type, x + 0.25, 1.88, 1.52, fill=SOFT, color=INK)
        add_text(slide, subtitle, x + 2.0, 1.93, 3.25, 0.27, size=11, color=MUTED, bold=True)
        add_text(slide, "USER", x + 0.25, 2.51, 0.8, 0.25, size=9, color=MUTED, bold=True)
        add_rect(slide, x + 0.25, 2.82, 5.22, 2.25, fill=SOFT, line=LINE)
        add_text(
            slide,
            prompt,
            x + 0.48,
            3.04,
            4.76,
            1.83,
            size=11.5,
            color=INK,
        )
        add_text(slide, "ASSISTANT", x + 0.25, 5.35, 0.95, 0.25, size=9, color=MUTED, bold=True)
        add_text(
            slide,
            "New topic: Courts & Legal Cases",
            x + 1.34,
            5.31,
            3.95,
            0.32,
            size=14,
            bold=True,
        )
    add_text(
        slide,
        "Real paired sample · source_sample_id: train:40:20",
        3.8,
        6.68,
        5.75,
        0.27,
        size=10.5,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Interest-only contains second-level interest names only; no first-level category is shown.")

    # 10 — Model training
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Model Training", "Compare supervised learning with reward-based optimization.", 10)
    training_tracks = [
        (
            "SFT",
            "Learn the task format and observed history-to-interest transitions from labeled examples.",
        ),
        (
            "GRPO",
            "Optimize generated predictions using task rewards without requiring a candidate-ranking head.",
        ),
        (
            "SFT → GRPO",
            "Initialize with supervised alignment, then refine generation using reward optimization.",
        ),
    ]
    for index, (title, detail) in enumerate(training_tracks):
        x = 0.72 + index * 4.08
        add_rect(slide, x, 1.68, 3.76, 2.15, fill=WHITE, line=LINE)
        add_text(slide, title, x + 0.22, 1.98, 3.3, 0.42, size=22, bold=True)
        add_text(slide, detail, x + 0.22, 2.62, 3.25, 0.82, size=12.5, color=MUTED)
    add_text(slide, "Training flow", 0.72, 4.28, 3.2, 0.34, size=16, bold=True)
    add_rect(slide, 0.72, 4.82, 5.72, 1.35, fill=WHITE, line=LINE)
    add_text(slide, "SFT", 1.0, 5.08, 1.0, 0.3, size=17, bold=True)
    add_text(
        slide,
        "Learn direct history-to-interest generation.",
        2.03,
        5.08,
        3.92,
        0.32,
        size=13,
    )
    add_text(
        slide,
        "Provides task alignment and initialization.",
        2.03,
        5.53,
        3.8,
        0.28,
        size=11.5,
        color=MUTED,
    )
    add_rect(slide, 6.89, 4.82, 5.72, 1.35, fill=WHITE, line=LINE)
    add_text(slide, "GRPO", 7.18, 5.08, 0.8, 0.3, size=17, bold=True)
    add_text(
        slide,
        "Always reason with CoT, then produce the final interest label.",
        8.0,
        5.08,
        4.18,
        0.42,
        size=13,
    )
    add_text(
        slide,
        "Reward optimization is applied to both reasoning and final prediction.",
        8.0,
        5.57,
        4.18,
        0.4,
        size=11.5,
        color=MUTED,
    )
    add_footer(slide, "Planned comparison across Behavior-only and Interest-only inputs.")

    # 11 — Results placeholder
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Experimental Results", "Recall and NDCG results will be added after model runs.", 11)
    columns = [
        ("Method", 0.82, 2.0),
        ("Input", 2.82, 2.05),
        ("Reasoning", 4.88, 1.75),
        ("Recall", 8.36, 1.3),
        ("NDCG", 10.36, 1.3),
    ]
    add_rect(slide, 0.72, 1.73, 12.0, 0.54, fill=SOFT, line=LINE)
    for label, x, w in columns:
        add_text(slide, label, x, 1.89, w, 0.23, size=10.5, bold=True)
    rows = [
        ("SFT", "Behavior-only", "Direct"),
        ("SFT", "Interest-only", "Direct"),
        ("GRPO", "Behavior-only", "CoT"),
        ("GRPO", "Interest-only", "CoT"),
        ("SFT → GRPO", "Both views", "CoT"),
    ]
    for index, (method, input_view, reasoning) in enumerate(rows):
        y = 2.27 + index * 0.57
        add_line(slide, 0.72, y + 0.54, 12.72, y + 0.54, color=LINE, width=0.7)
        add_text(slide, method, 0.82, y + 0.14, 1.8, 0.24, size=11.5, bold=True)
        add_text(slide, input_view, 2.82, y + 0.14, 1.9, 0.24, size=11.5)
        add_text(slide, reasoning, 4.88, y + 0.14, 2.1, 0.24, size=11.5)
        add_text(slide, "—", 8.36, y + 0.1, 1.3, 0.3, size=15, align=PP_ALIGN.CENTER)
        add_text(slide, "—", 10.36, y + 0.1, 1.3, 0.3, size=15, align=PP_ALIGN.CENTER)
    add_rect(slide, 0.72, 5.58, 5.72, 1.0, fill=WHITE, line=LINE)
    add_text(slide, "Recall", 0.98, 5.87, 1.1, 0.3, size=17, bold=True)
    add_text(slide, "Chart placeholder", 3.82, 5.89, 2.15, 0.25, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    add_rect(slide, 6.89, 5.58, 5.72, 1.0, fill=WHITE, line=LINE)
    add_text(slide, "NDCG", 7.15, 5.87, 1.1, 0.3, size=17, bold=True)
    add_text(slide, "Chart placeholder", 9.99, 5.89, 2.15, 0.25, size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    add_footer(slide, "Results pending. Metric cutoffs and judge configuration will be reported with the experiments.")

    # 12 — Related work landscape
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "Related Work", "Four research streams converge—but leave a gap.", 12)
    cards = [
        (
            "01",
            "Sequential\nrecommendation",
            "Predict the next item from ordered behavior.",
            "Kang & McAuley, 2018\nSun et al., 2019",
            INDIGO_PALE,
            INDIGO,
        ),
        (
            "02",
            "Multi-interest\nmodeling",
            "Represent diverse known preferences with multiple vectors.",
            "Li et al., 2019\nCen et al., 2020",
            SOFT,
            MINT_DARK,
        ),
        (
            "03",
            "Diversity &\nexploration",
            "Move beyond dominant preferences and redundant lists.",
            "Li et al., 2010\nGe et al., 2010",
            GOLD_PALE,
            INK,
        ),
        (
            "04",
            "Generative\nrecommendation",
            "Produce items or semantic outputs with language models.",
            "Geng et al., 2022\nWang et al., 2024",
            CORAL_PALE,
            CORAL,
        ),
    ]
    for index, (num, title, detail, refs, fill, accent) in enumerate(cards):
        x = 0.72 + index * 3.02
        add_rect(slide, x, 1.72, 2.76, 4.58, fill=WHITE, line=LINE)
        add_circle(slide, num, x + 0.24, 1.98, 0.52, fill=fill, color=accent, size=11)
        add_text(slide, title, x + 0.24, 2.72, 2.25, 0.9, size=18, bold=True)
        add_text(slide, detail, x + 0.24, 3.78, 2.22, 1.0, size=12.5, color=MUTED)
        add_line(slide, x + 0.24, 5.03, x + 2.5, 5.03, color=LINE, width=1)
        add_text(slide, refs, x + 0.24, 5.25, 2.28, 0.62, size=9.5, color=accent, bold=True)
    add_rect(slide, 3.28, 6.52, 6.78, 0.48, fill=SOFT, line=LINE)
    add_text(
        slide,
        "Missing link: open-ended prediction of an unseen semantic interest",
        3.48,
        6.64,
        6.38,
        0.22,
        size=11.5,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, "Representative work; full citations appear on the final slide.")

    # 13 — Closest prior work
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(
        slide,
        "Closest Prior Work",
        "Wang et al. (2024): LLMs for User Interest Exploration",
        13,
    )
    add_rect(slide, 0.72, 1.58, 6.43, 4.98, fill=WHITE, line=LINE)
    if PRIOR_WORK_IMAGE.is_file():
        slide.shapes.add_picture(
            str(PRIOR_WORK_IMAGE),
            Inches(0.94),
            Inches(1.86),
            width=Inches(5.98),
        )
    add_text(
        slide,
        "Figure 2: SFT prompt using K = 2 recent interest clusters",
        1.01,
        6.18,
        5.82,
        0.28,
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    add_rect(slide, 7.48, 1.58, 5.14, 4.98, fill=WHITE, line=LINE)
    add_text(slide, "Paper at a glance", 7.82, 1.86, 2.2, 0.3, size=11, color=INK, bold=True)
    paper_summary = [
        ("Problem", "Feedback loops reinforce past consumption and limit interest discovery."),
        ("Approach", "An LLM generates a novel interest cluster; a sequential model retrieves personalized items within it."),
        ("Alignment", "SFT on diverse real-world transitions produces controlled outputs over 761 predefined clusters."),
        ("Deployment", "K = 2 cluster-pair predictions are precomputed offline and served through table lookup."),
        ("Evidence", "Live experiments report stronger novelty, broader interest exploration, and user-growth gains."),
    ]
    for index, (label, value) in enumerate(paper_summary):
        y = 2.3 + index * 0.72
        add_text(slide, label.upper(), 7.82, y, 1.0, 0.25, size=8.5, color=MUTED, bold=True)
        add_text(slide, value, 8.88, y - 0.03, 3.35, 0.55, size=11.5, color=INK, bold=True)
    add_text(
        slide,
        "Why it matters here: it directly motivates using LLMs to predict novel interests; our study isolates this prediction layer for news with a fixed 125-label taxonomy.",
        7.82,
        5.91,
        4.25,
        0.45,
        size=10.5,
        color=INK,
        bold=True,
    )
    add_footer(
        slide,
        "Wang et al. (2024), LLMs for User Interest Exploration in Large-scale Recommendation Systems, RecSys. arXiv:2405.16363 · DOI: 10.1145/3640457.3688161",
    )

    # 14 — Research gap and questions
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_text(slide, "RESEARCH GAP & QUESTIONS", 0.78, 0.52, 3.4, 0.3, size=10, color=MINT, bold=True)
    add_text(
        slide,
        "14",
        12.12,
        0.47,
        0.5,
        0.3,
        size=10,
        color=MUTED,
        bold=True,
        align=PP_ALIGN.RIGHT,
    )
    add_text(slide, "A clear gap leads to three research questions.", 0.78, 0.94, 9.5, 0.62, size=31, color=INK, bold=True)
    add_rect(slide, 0.8, 1.59, 11.72, 0.64, fill=SOFT, line=LINE)
    add_text(
        slide,
        "Can generative models infer the semantic direction of preference expansion—not merely the next item?",
        1.06,
        1.76,
        11.2,
        0.28,
        size=13.5,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    rqs = [
        (
            "RQ1",
            "Generative predictability",
            "Can a language model predict the next novel fine-grained interest without a predefined candidate set?",
            SOFT,
            MINT_DARK,
        ),
        (
            "RQ2",
            "History representation",
            "How does Behavior-only compare with Interest-only for novel-interest prediction?",
            CORAL_PALE,
            CORAL,
        ),
        (
            "RQ3",
            "Learning & generalization",
            "Does performance reflect personalized interest transitions rather than global popularity?",
            INDIGO_PALE,
            INDIGO,
        ),
    ]
    for index, (rq, title, question, fill, accent) in enumerate(rqs):
        y = 2.55 + index * 1.28
        add_circle(slide, rq, 0.82, y, 0.72, fill=fill, color=accent, size=12)
        add_text(slide, title, 1.83, y + 0.02, 3.35, 0.3, size=15, color=MUTED, bold=True)
        add_text(slide, question, 1.83, y + 0.39, 9.95, 0.54, size=16, color=INK, bold=True)
        if index < 2:
            add_line(slide, 1.82, y + 1.02, 12.05, y + 1.02, color=LINE, width=1)
    add_text(
        slide,
        "Next: controlled experiments across training paradigms, data scales, and reasoning settings.",
        0.82,
        6.75,
        11.6,
        0.32,
        size=11,
        color=MUTED,
        bold=True,
    )

    # 15 — References
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_slide_header(slide, "References", "Verified foundations for the research framing.", 15)
    refs_left = [
        "Kang, W.-C., & McAuley, J. (2018). Self-Attentive Sequential Recommendation. ICDM. DOI: 10.1109/ICDM.2018.00035",
        "Sun, F. et al. (2019). BERT4Rec. CIKM. DOI: 10.1145/3357384.3357895",
        "Li, C. et al. (2019). Multi-Interest Network with Dynamic Routing for Recommendation at Tmall. CIKM. DOI: 10.1145/3357384.3357814",
        "Cen, Y. et al. (2020). Controllable Multi-Interest Framework for Recommendation. KDD. DOI: 10.1145/3394486.3403344",
        "Li, L. et al. (2010). A Contextual-Bandit Approach to Personalized News Article Recommendation. WWW. DOI: 10.1145/1772690.1772758",
    ]
    refs_right = [
        "Ge, M. et al. (2010). Beyond Accuracy: Evaluating Recommender Systems by Coverage and Serendipity. RecSys. DOI: 10.1145/1864708.1864761",
        "Geng, S. et al. (2022). Recommendation as Language Processing (P5). RecSys. DOI: 10.1145/3523227.3546767",
        "Bao, K. et al. (2023). TALLRec: Aligning Large Language Models with Recommendation. RecSys. DOI: 10.1145/3604915.3608857",
        "Wang, J. et al. (2024). LLMs for User Interest Exploration in Large-scale Recommendation Systems. RecSys. DOI: 10.1145/3640457.3688161",
    ]
    for col, refs in enumerate((refs_left, refs_right)):
        x = 0.72 + col * 6.18
        for index, ref in enumerate(refs):
            y = 1.68 + index * 0.96
            add_circle(slide, str(index + 1 + col * 5), x, y + 0.02, 0.4, fill=INDIGO_PALE, color=INDIGO, size=9)
            add_text(slide, ref, x + 0.58, y, 5.15, 0.78, size=10.5, color=INK)
    add_rect(slide, 0.72, 6.5, 12.0, 0.42, fill=SOFT, line=LINE)
    add_text(
        slide,
        "All citations were verified against DOI records or the official paper.",
        0.95,
        6.61,
        11.55,
        0.2,
        size=10,
        color=MINT_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide)

    # Core properties.
    prs.core_properties.title = "Generative Novel Interest Prediction from User History"
    prs.core_properties.subject = "Research framing and questions"
    prs.core_properties.author = "River"
    prs.core_properties.keywords = "recommendation, novel interest, generative recommendation, MIND"

    return prs


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation = make_deck()
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
